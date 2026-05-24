#!/usr/bin/env python3
"""
立高门业报价单生成器 — Excel → PDF
用法: python3 quote_xlsx.py quote-data.json
输出: 报价单-{客户}-{日期}.xlsx + .pdf
"""

import json, sys, os, io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XlImage

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SEAL_PATH = os.path.join(SCRIPT_DIR, "images/seal.png")

# 颜色常量
DARK_BLUE  = "1a3a5c"
WHITE      = "FFFFFF"
LIGHT_GRAY = "F5F5F5"
RED_ACCENT = "D4380D"
BORDER_GRAY = "CCCCCC"

# 通用边框
thin_border = Border(
    left=Side(style='thin', color=BORDER_GRAY),
    right=Side(style='thin', color=BORDER_GRAY),
    top=Side(style='thin', color=BORDER_GRAY),
    bottom=Side(style='thin', color=BORDER_GRAY),
)
header_fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
header_font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
title_font  = Font(name="微软雅黑", size=20, bold=True, color=WHITE)
info_font   = Font(name="微软雅黑", size=9, color=WHITE)
data_font   = Font(name="微软雅黑", size=10)
bold_font   = Font(name="微软雅黑", size=10, bold=True)
total_font  = Font(name="微软雅黑", size=12, bold=True, color=RED_ACCENT)
note_font   = Font(name="微软雅黑", size=9, color="888888")
gray_fill   = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type='solid')
center   = Alignment(horizontal='center', vertical='center', wrap_text=True)
left     = Alignment(horizontal='left', vertical='center', wrap_text=True)
right    = Alignment(horizontal='right', vertical='center')
merge_c  = Alignment(horizontal='center', vertical='center')


def num_to_chinese(n):
    digits = "零壹贰叁肆伍陆柒捌玖"
    units = ["", "拾", "佰", "仟"]
    if n == 0: return "零元整"
    if n >= 100000000: return f"{n:,}元整"
    wan = n // 10000; rest = n % 10000
    def conv4(num):
        if num == 0: return "零"
        s = f"{num:04d}"; r, pz = "", False
        for i, ch in enumerate(s):
            d = int(ch); pos = 3 - i
            if d == 0: pz = True
            else:
                if pz and r: r += "零"
                r += digits[d] + units[pos]; pz = False
        return r.rstrip("零")
    result = ""
    if wan > 0: result += conv4(wan) + "万"
    if rest > 0:
        if wan > 0 and rest < 1000: result += "零"
        result += conv4(rest)
    return result + "元整"


def build_xlsx(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"

    # 列宽
    col_widths = [5, 24, 10, 10, 8, 9, 12, 12, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 页边距 / 打印设置
    ws.sheet_properties.pageSetUpPr = None
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9  # A4
    ws.page_margins.left = 0.5; ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5; ws.page_margins.bottom = 0.5

    row = 1

    # ══════ 标题栏 ══════
    ws.merge_cells(start_row=row, start_column=1, end_row=row+1, end_column=7)
    c = ws.cell(row=row, column=1, value="立高门业报价单")
    c.font = title_font; c.fill = header_fill; c.alignment = Alignment(horizontal='left', vertical='center')
    for col in range(1, 8):
        ws.cell(row=row, column=col).fill = header_fill
        ws.cell(row=row+1, column=col).fill = header_fill
    # 公司信息右侧
    ws.merge_cells(start_row=row, start_column=8, end_row=row+1, end_column=9)
    info_text = "重庆立高门业有限公司\n地址：巴南区公平场大道845号附17号\n电话：023-67399080 | 郭纪敏 18983919000"
    c = ws.cell(row=row, column=8, value=info_text)
    c.font = info_font; c.fill = header_fill; c.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    for col in range(8, 10):
        ws.cell(row=row, column=col).fill = header_fill
        ws.cell(row=row+1, column=col).fill = header_fill
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row+1].height = 22
    row += 2

    # ══════ 客户信息 ══════
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cust_str = f"客户：{data.get('customer','')}　　联系人：{data.get('contact','')}　　电话：{data.get('phone','')}　　报价日期：{data.get('date','')}"
    c = ws.cell(row=row, column=1, value=cust_str)
    c.font = Font(name="微软雅黑", size=11); c.fill = gray_fill; c.alignment = Alignment(horizontal='left', vertical='center')
    for col in range(1, 10):
        ws.cell(row=row, column=col).fill = gray_fill
    ws.row_dimensions[row].height = 22
    row += 1

    # ══════ 表头 ══════
    headers = ["序号", "产品名称", "门洞宽\n(mm)", "门洞高\n(mm)", "数量\n(樘)", "面积\n(㎡)", "单价\n(元/樘)", "金额\n(元)", "备注"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = center; c.border = thin_border
    ws.row_dimensions[row].height = 28
    row += 1

    # ══════ 数据行 ══════
    total = 0
    items = data.get("items", [])
    for idx, item in enumerate(items):
        no = item.get("no", idx+1)
        name = item.get("name", "")
        w_val = item.get("width", "-")
        h_val = item.get("height", "-")
        qty = item.get("qty", "")
        area = item.get("area", "-")
        up = item.get("unit_price", "")
        amt = item.get("amount", 0) or (int(qty) * int(up) if qty and up else 0)
        notes = item.get("notes", "")
        total += amt

        vals = [no, name, w_val, h_val, qty, area,
                up if isinstance(up, str) or up == "" else up,
                amt, notes]
        row_fill = gray_fill if idx % 2 == 0 else None
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = data_font; c.border = thin_border
            c.alignment = left if i in (2, 9) else (right if i in (7, 8) else center)
            if row_fill:
                c.fill = row_fill
            # 金额格式化
            if i == 8 and isinstance(amt, int) and amt > 0:
                c.number_format = '#,##0'
            if i == 7 and isinstance(up, int) and up > 0:
                c.number_format = '#,##0'
        ws.row_dimensions[row].height = 20
        row += 1

    # ══════ 合计行 ══════
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    c = ws.cell(row=row, column=1, value="合  计（含13%增值税）")
    c.font = Font(name="微软雅黑", size=12, bold=True); c.fill = PatternFill(start_color="E8F0FA", end_color="E8F0FA", fill_type='solid')
    c.alignment = Alignment(horizontal='right', vertical='center')
    for col in range(1, 10):
        ws.cell(row=row, column=col).fill = PatternFill(start_color="E8F0FA", end_color="E8F0FA", fill_type='solid')
        ws.cell(row=row, column=col).border = Border(top=Side(style='medium', color=DARK_BLUE), bottom=Side(style='thin', color=BORDER_GRAY))
    c2 = ws.cell(row=row, column=8, value=total)
    c2.font = total_font; c2.alignment = right; c2.number_format = '#,##0'
    c2.border = Border(top=Side(style='medium', color=DARK_BLUE), bottom=Side(style='thin', color=BORDER_GRAY), left=Side(style='thin', color=BORDER_GRAY), right=Side(style='thin', color=BORDER_GRAY))
    ws.row_dimensions[row].height = 24
    row += 2

    # ══════ 大写金额 ══════
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    chinese = num_to_chinese(total)
    c = ws.cell(row=row, column=1, value=f"合计（大写）：{chinese}")
    c.font = Font(name="微软雅黑", size=12, bold=True, color=RED_ACCENT)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 28
    row += 1

    # ══════ 电子章 ══════
    if os.path.exists(SEAL_PATH):
        img = XlImage(SEAL_PATH)
        img.width = 80; img.height = 80
        img.anchor = ws.cell(row=row, column=8).coordinate  # type: ignore
        ws.add_image(img)
        c = ws.cell(row=row, column=8, value="报价专用章")
        c.font = Font(name="微软雅黑", size=8, color="999999"); c.alignment = center
    ws.row_dimensions[row].height = 80
    row = max(row + 6, row + 1)

    # ══════ 备注 ══════
    row += 1
    notes = [
        "备注：",
        "1、以上报价含整套门制作、运输、安装等全部费用；",
        "2、报价不含拆旧门、拆墙、加装门框、门控制箱以上强电布线；",
        "3、以上报价含13%增值税专用发票，本报价有效期30天。",
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        c = ws.cell(row=row, column=1, value=note)
        c.font = note_font; c.alignment = Alignment(horizontal='left', vertical='center')
        row += 1

    # ══════ 页脚 ══════
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    c = ws.cell(row=row, column=1, value="重庆立高门业有限公司 · 15年工业门自研品牌 · 023-67399080")
    c.font = Font(name="微软雅黑", size=8, color="BBBBBB"); c.alignment = center

    # 打印区域
    ws.print_area = f"A1:I{row}"

    return wb


def main():
    if len(sys.argv) < 2:
        print("用法: python3 quote_xlsx.py <quote-data.json>")
        sys.exit(1)

    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = build_xlsx(data)

    customer_short = data.get("customer", "客户")[:6].replace(" ", "")
    date_str = data.get("date", datetime.now().strftime("%Y%m%d")).replace("-", "")
    base_name = f"报价单-{customer_short}-{date_str}"

    xlsx_path = os.path.join(SCRIPT_DIR, f"{base_name}.xlsx")
    wb.save(xlsx_path)
    print(f"✅ Excel 已生成: {xlsx_path}")

    # Excel → PDF via Windows Excel COM (WSL环境)
    pdf_path = os.path.join(SCRIPT_DIR, f"{base_name}.pdf")
    win_xlsx = f"C:\\\\temp\\\\quote_temp.xlsx"
    win_pdf  = f"C:\\\\temp\\\\quote_temp.pdf"

    # 复制到Windows可访问路径
    import shutil, subprocess
    shutil.copy(xlsx_path, "/mnt/c/temp/quote_temp.xlsx")

    ps_cmd = (
        f"$xlsx = '{win_xlsx}';"
        f"$excel = New-Object -ComObject Excel.Application;"
        f"$excel.Visible = $false; $excel.DisplayAlerts = $false;"
        f"$wb = $excel.Workbooks.Open($xlsx);"
        f"$wb.ExportAsFixedFormat(0, '{win_pdf}');"
        f"$wb.Close(); $excel.Quit();"
        f"[System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null"
    )
    ret = subprocess.run(["powershell.exe", "-Command", ps_cmd],
                         capture_output=True, text=True, timeout=30)
    if ret.returncode == 0 and os.path.exists("/mnt/c/temp/quote_temp.pdf"):
        shutil.copy("/mnt/c/temp/quote_temp.pdf", pdf_path)
        print(f"✅ PDF 已生成: {pdf_path}")
    else:
        print(f"⚠️  PDF转换失败，Excel文件可用Excel打开后另存为PDF")
        print(f"   {ret.stderr[:200] if ret.stderr else ''}")

    total = sum(i.get("amount", 0) or (int(i.get("qty",0))*int(i.get("unit_price",0)) if i.get("qty") and i.get("unit_price") else 0) for i in data.get("items",[]))
    print(f"   客户: {data.get('customer')}")
    print(f"   总金额: ¥{total:,}")


if __name__ == "__main__":
    main()
